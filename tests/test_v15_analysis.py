"""v1.5.0 analysis tools: site config, pure helpers, profile, brand split,
striking distance, movers, cannibalisation, recrawl worklist, the safe
fetcher and sitemap diff, and the UI-export SQL sessions."""
from __future__ import annotations

import asyncio
import json
import sqlite3
import zipfile
from unittest.mock import MagicMock

import pytest

import gsc_server
from gsc_server import (
    gsc_brand_split,
    gsc_cannibalisation,
    gsc_load_ui_export,
    gsc_movers,
    gsc_page_query_profile,
    gsc_query_ui_export,
    gsc_recrawl_worklist,
    gsc_sitemap_diff,
    gsc_striking_distance,
)

SITE = "sc-domain:example.com"


def _row(keys, clicks, impressions, ctr=0.0, position=5.0):
    keys = keys if isinstance(keys, list) else [keys]
    return {"keys": keys, "clicks": clicks, "impressions": impressions, "ctr": ctr, "position": position}


def _service(respond):
    service = MagicMock()
    service.bodies = []

    def _query(*, siteUrl, body):
        service.bodies.append(json.loads(json.dumps(body)))
        req = MagicMock()
        req.execute.return_value = respond(body)
        return req
    service.searchanalytics.return_value.query.side_effect = _query
    return service


def _patch(monkeypatch, service):
    monkeypatch.setattr(gsc_server, "get_gsc_service", lambda: service)


def _filters(body):
    return [f for g in body.get("dimensionFilterGroups", []) for f in g["filters"]]


def _write_config(cfg):
    with open(gsc_server._site_config_path(), "w") as fh:
        json.dump(cfg, fh)


# ---------------------------------------------------------------------------
# Pure helpers
# ---------------------------------------------------------------------------


class TestHelpers:
    @pytest.mark.parametrize("q,lang", [
        ("cuentas incobrables", "es"), ("cómo cobrar una factura", "es"),
        ("créances irrécouvrables recouvrement", "fr"), ("rechnung zahlung mahnung", "de"),
        ("uncollectible accounts", "en"), ("счета к оплате", "other"),
    ])
    def test_language_guess(self, q, lang):
        assert gsc_server._guess_language(q)[0] == lang

    def test_machine_flags(self):
        rows = [
            {"query": "sales cash collection performance", "clicks": 0, "impressions": 600},
            {"query": "+site.example.com pricing", "clicks": 0, "impressions": 10},
            {"query": "site:example.com", "clicks": 1, "impressions": 5},
            {"query": "how do i write a polite email asking a client to pay an overdue invoice", "clicks": 0, "impressions": 3},
            {"query": "spiky", "clicks": 1, "impressions": 200},
        ]
        daily = {"spiky": [2, 1, 3, 180, 2, 1, 2]}
        flags = gsc_server._machine_query_flags(rows, 1000, daily)
        assert flags["sales cash collection performance"] == ["dominant_zero_click"]
        assert "operator_string" in flags["+site.example.com pricing"]
        assert "operator_string" in flags["site:example.com"]
        assert "assistant_style" in flags["how do i write a polite email asking a client to pay an overdue invoice"]
        assert flags["spiky"] == ["spike"]

    def test_position_buckets_and_ctr_curve(self):
        rows = [
            {"position": 2, "impressions": 100, "ctr": 0.3},
            {"position": 2.4, "impressions": 100, "ctr": 0.1},
            {"position": 7, "impressions": 200, "ctr": 0.05},
            {"position": 15, "impressions": 100, "ctr": 0.01},
        ]
        b = gsc_server._position_buckets(rows)
        assert b == {"1-3": 0.4, "4-10": 0.4, "11-20": 0.2, "21+": 0.0}
        curve = gsc_server._site_ctr_curve(rows)
        assert curve[2] == pytest.approx(0.2)  # median of 0.3 and 0.1
        assert gsc_server._expected_ctr(8, curve) == 0.05  # nearest known position

    def test_terms_regex_joins_and_validates(self):
        assert gsc_server._terms_regex(["a", "b"], field="brand") == "(?:a)|(?:b)"
        with pytest.raises(gsc_server._SaValidationError):
            gsc_server._terms_regex(["(?!x)"], field="brand")

    def test_date_buckets(self):
        assert gsc_server._date_bucket("2026-09-28", "month") == "2026-09"
        assert gsc_server._date_bucket("2026-09-28", "week") == "2026-W40"
        assert gsc_server._date_bucket("2026-09-28", "day") == "2026-09-28"


# ---------------------------------------------------------------------------
# Brand split (acceptance test 7's mechanism)
# ---------------------------------------------------------------------------


class TestBrandSplit:
    async def test_split_uses_api_regex_and_config_defaults(self, monkeypatch):
        _write_config({SITE: {"brand_terms": ["example"], "login_terms": ["^(example login)$"]}})

        def respond(body):
            f = _filters(body)
            base = {"2026-08-01": 100, "2026-08-02": 50, "2026-09-01": 80}
            if not f:
                return {"rows": [_row(d, c, c * 10) for d, c in base.items()]}
            op = f[0]["operator"]
            share = {"includingRegex": 0.2, "excludingRegex": 0.7}[op]
            if f[0]["expression"].startswith("^"):
                share = 0.05
            return {"rows": [_row(d, int(c * share), int(c * share * 10)) for d, c in base.items()]}
        service = _service(respond)
        _patch(monkeypatch, service)
        out = await gsc_brand_split(SITE, "2026-08-01", "2026-09-01")
        assert out["ok"] is True
        aug = next(r for r in out["rows"] if r["period"] == "2026-08")
        assert (aug["total_clicks"], aug["branded_clicks"], aug["non_branded_clicks"], aug["login_clicks"]) == (150, 30, 105, 7)
        assert aug["anonymised_clicks"] == 15
        assert aug["brand_share_clicks"] == pytest.approx(0.2)
        assert aug["days_of_data"] == 2
        ops = sorted((f["operator"], f["expression"]) for b in service.bodies for f in _filters(b))
        assert ("excludingRegex", "example") in ops and ("includingRegex", "^(example login)$") in ops

    async def test_needs_brand_terms(self, monkeypatch):
        _patch(monkeypatch, _service(lambda b: {"rows": []}))
        out = await gsc_brand_split(SITE)
        assert out["error_code"] == "BAD_REQUEST" and "brand" in out["error"].lower()


# ---------------------------------------------------------------------------
# Page profile, striking distance, movers, cannibalisation
# ---------------------------------------------------------------------------


class TestProfile:
    async def test_profile_sorts_flags_and_shares(self, monkeypatch):
        _write_config({SITE: {"brand_terms": ["example"]}})

        def respond(body):
            dims = body["dimensions"]
            if dims == []:
                return {"rows": [_row([], 12, 2000)]}
            if dims == ["query", "date"]:
                return {"rows": [_row(["cuentas incobrables", "2026-09-01"], 0, 5)]}
            return {"rows": [
                _row("example login", 10, 100, position=1.5),
                _row("cuentas incobrables", 0, 1200, position=8),
                _row("site:example.com", 0, 50, position=30),
            ]}
        _patch(monkeypatch, _service(respond))
        out = await gsc_page_query_profile(SITE, "https://example.com/p")
        assert out["ok"] is True
        assert [q["impressions"] for q in out["queries"]] == [1200, 100, 50]
        s = out["summary"]
        assert s["page_total"] == {"clicks": 12, "impressions": 2000}
        assert s["non_english_languages"] == ["es"]
        assert s["non_english_share"]["of_page_total_impressions"] == pytest.approx(0.6)
        assert s["brand_share"]["of_query_rows_clicks"] == pytest.approx(1.0)
        assert s["flag_counts"]["dominant_zero_click"] == 1 and s["flag_counts"]["operator_string"] == 1
        assert s["position_distribution"]["4-10"] == pytest.approx(1200 / 1350)


class TestStrikingDistance:
    async def test_candidates_ranked_by_upside_with_top_queries(self, monkeypatch):
        def respond(body):
            if body["dimensions"] == ["page"]:
                return {"rows": [
                    _row("https://example.com/a", 5, 1000, ctr=0.005, position=8),
                    _row("https://example.com/b", 40, 1000, ctr=0.04, position=8),
                    _row("https://example.com/c", 1, 5000, ctr=0.0002, position=12),
                    _row("https://example.com/skip", 0, 9000, ctr=0.0, position=9),
                    _row("https://example.com/top", 90, 1000, ctr=0.09, position=2),
                ]}
            return {"rows": [_row(["https://example.com/c", "kw c"], 1, 4000)]}
        _patch(monkeypatch, _service(respond))
        _write_config({SITE: {"ctr_curve": {"8": 0.03, "12": 0.01, "2": 0.1}}})
        out = await gsc_striking_distance(SITE, exclude_urls=["https://example.com/skip"])
        pages = [c["page"] for c in out["candidates"]]
        assert pages == ["https://example.com/c", "https://example.com/a"]  # b is above expected; top out of band
        assert out["candidates"][0]["click_upside"] == pytest.approx((0.01 - 0.0002) * 5000)
        assert out["candidates"][0]["top_queries"][0]["query"] == "kw c"
        assert out["meta"]["ctr_curve_source"] == "site_config"


class TestMovers:
    async def test_gainers_losers_and_share(self, monkeypatch):
        calls = {"n": 0}

        def respond(body):
            calls["n"] += 1
            if calls["n"] == 1:
                return {"rows": [_row("/a", 10, 100), _row("/b", 50, 500), _row("/gone", 5, 50)]}
            return {"rows": [_row("/a", 40, 300), _row("/b", 20, 400), _row("/new", 5, 50)]}
        _patch(monkeypatch, _service(respond))
        out = await gsc_movers(SITE, days=7)
        assert out["net_change"]["clicks"] == 0  # +30 -30 +5 -5
        assert [g["page"] for g in out["gainers"]] == ["/a", "/new"]
        assert [l["page"] for l in out["losers"]] == ["/b", "/gone"]
        assert out["meta"]["period_b"]["end_date"] == "2026-09-25"


class TestCannibalisation:
    async def test_groups_queries_and_collapses_fragments(self, monkeypatch):
        _patch(monkeypatch, _service(lambda b: {"rows": [
            _row(["sage 200", "https://example.com/x"], 3, 100, position=4),
            _row(["sage 200", "https://example.com/y"], 1, 60, position=9),
            _row(["sage versions", "https://example.com/z"], 1, 50),
            _row(["sage versions", "https://example.com/z#part"], 0, 40),
        ]}))
        out = await gsc_cannibalisation(SITE, query_regex="\\bsage\\b")
        assert [q["query"] for q in out["queries"]] == ["sage 200"]
        assert out["queries"][0]["pages"][0]["impression_share"] == pytest.approx(100 / 160)


# ---------------------------------------------------------------------------
# Recrawl worklist (acceptance test 6's worklist half)
# ---------------------------------------------------------------------------


def _inspection(verdict, coverage, crawl, g=None, u=None):
    idx = {"verdict": verdict, "coverageState": coverage, "lastCrawlTime": crawl}
    if g:
        idx.update(googleCanonical=g, userCanonical=u)
    return {"inspectionResult": {"indexStatusResult": idx}}


class TestRecrawlWorklist:
    def _setup(self, monkeypatch, by_url):
        service = MagicMock()

        def _inspect(body):
            req = MagicMock()
            req.execute.return_value = by_url[body["inspectionUrl"]]
            return req
        service.urlInspection.return_value.index.return_value.inspect.side_effect = _inspect
        service.searchanalytics.return_value.query.return_value.execute.return_value = {
            "rows": [_row("https://example.com/old", 9, 90)]}
        monkeypatch.setattr(gsc_server, "get_gsc_service", lambda: service)

        async def fake_statuses(client, urls, site_url, **kw):
            return {u: {"url": u, "status": 200} for u in urls}

        async def fake_control(client, site_url, sample):
            return {"url": "https://example.com/", "ok": True, "status": 200}
        monkeypatch.setattr(gsc_server, "_fetch_statuses", fake_statuses)
        monkeypatch.setattr(gsc_server, "_control_ok", fake_control)

    async def test_pending_then_complete_worklist(self, monkeypatch):
        by_url = {
            "https://example.com/old": _inspection("PASS", "Submitted and indexed", "2026-09-10T00:00:00Z"),
            "https://example.com/fresh": _inspection("PASS", "Submitted and indexed", "2026-09-26T00:00:00Z"),
            "https://example.com/redir": _inspection("NEUTRAL", "Page with redirect", "2026-09-20T00:00:00Z"),
            "https://example.com/dup": _inspection("NEUTRAL", "Duplicate, Google chose different canonical than user",
                                                   "2026-09-25T00:00:00Z", "https://example.com/other", "https://example.com/dup"),
            "https://example.com/noidx": _inspection("NEUTRAL", "Crawled - currently not indexed", "2026-09-24T00:00:00Z"),
        }
        self._setup(monkeypatch, by_url)
        items = [{"url": u, "changed_on": "2026-09-14"} for u in by_url]
        first = await gsc_recrawl_worklist(SITE, items)
        assert first["status"] == "pending" and first["poll_with"] == "gsc_inspect_status"
        for _ in range(200):
            out = await gsc_recrawl_worklist(SITE, items)
            if out["status"] == "complete":
                break
            await asyncio.sleep(0.01)
        assert out["counts"] == {"needs_request_indexing": 3, "canonical_mismatches": 1,
                                 "no_action_needed": 1, "unresolved": 0}
        needs = {e["url"]: e for e in out["needs_request_indexing"]}
        assert "crawled before the change" in needs["https://example.com/old"]["reasons"][0]
        assert "answers 200" in needs["https://example.com/redir"]["reasons"][0]
        assert needs["https://example.com/noidx"]["reasons"][0].startswith("not indexed")
        assert out["needs_request_indexing"][0]["url"] == "https://example.com/old"  # most clicks first
        assert out["canonical_mismatches"][0]["google_canonical"] == "https://example.com/other"

    async def test_daily_cap_and_failed_urls_are_unresolved(self, monkeypatch):
        by_url = {f"https://example.com/p{i}": _inspection("PASS", "ok", "2026-01-01T00:00:00Z") for i in range(3)}
        self._setup(monkeypatch, by_url)
        items = [{"url": u, "changed_on": "2026-09-14"} for u in by_url] + [{"url": "https://example.com/boom", "changed_on": "2026-09-14"}]
        orig = gsc_server._inspect_one

        async def flaky(service, site, url, **kw):
            if url.endswith("boom"):
                raise RuntimeError("inspect failed")
            return await orig(service, site, url, **kw)
        monkeypatch.setattr(gsc_server, "_inspect_one", flaky)
        await gsc_recrawl_worklist(SITE, items, daily_cap=2)
        for _ in range(200):
            out = await gsc_recrawl_worklist(SITE, items, daily_cap=2)
            if out["status"] == "complete":
                break
            await asyncio.sleep(0.01)
        assert len(out["needs_request_indexing"]) == 2 and len(out["needs_request_indexing_later"]) == 1, out
        assert out["paste_ready"].count("\n") == 1
        assert out["unresolved"][0]["url"] == "https://example.com/boom", out

    async def test_requires_change_dates(self, monkeypatch):
        self._setup(monkeypatch, {})
        out = await gsc_recrawl_worklist(SITE, ["https://example.com/a"])
        assert out["error_code"] == "BAD_REQUEST"


# ---------------------------------------------------------------------------
# Safe fetcher and sitemap diff
# ---------------------------------------------------------------------------


class TestSafeFetch:
    @pytest.mark.parametrize("url,site", [
        ("ftp://example.com/x", SITE),
        ("https://evil.test/x", SITE),
        ("https://example.com.evil.test/x", SITE),
        ("https://other.example.com/x", "https://www.example.com/"),
    ])
    async def test_out_of_scope_refused_before_any_request(self, url, site):
        client = MagicMock()
        with pytest.raises(gsc_server._FetchRefused):
            await gsc_server._safe_fetch(client, url, site)
        client.stream.assert_not_called()

    async def test_private_address_refused_before_request(self, monkeypatch):
        monkeypatch.setattr(gsc_server.socket, "getaddrinfo",
                            lambda *a, **k: [(None, None, None, None, ("127.0.0.1", 443))])
        client = MagicMock()
        with pytest.raises(gsc_server._FetchRefused):
            await gsc_server._safe_fetch(client, "https://www.example.com/", SITE)
        client.stream.assert_not_called()

    def test_subdomains_allowed_for_domain_property(self):
        assert gsc_server._host_allowed("www.example.com", SITE)
        assert gsc_server._host_allowed("example.com", SITE)
        assert not gsc_server._host_allowed("badexample.com", SITE)


class TestSitemapDiff:
    async def test_nested_index_image_locs_and_statuses(self, monkeypatch):
        index = b"""<?xml version="1.0"?><sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
          <sitemap><loc>https://www.example.com/sm-1.xml</loc></sitemap>
          <sitemap><loc>https://www.example.com/sitemap.xml</loc></sitemap></sitemapindex>"""
        child = b"""<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"
            xmlns:image="http://www.google.com/schemas/sitemap-image/1.1">
          <url><loc>https://www.example.com/a</loc><image:image><image:loc>https://cdn.test/i.png</image:loc></image:image></url>
          <url><loc>https://www.example.com/moved</loc></url>
          <url><loc>https://www.example.com/dead</loc></url></urlset>"""
        pages = {"https://www.example.com/": 200, "https://www.example.com/a": 200,
                 "https://www.example.com/moved": 301, "https://www.example.com/dead": 404,
                 "https://www.example.com/live-only": 200}

        async def fake_fetch(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": index}
            if url.endswith("sm-1.xml"):
                return {"url": url, "status": 200, "location": None, "body": child}
            return {"url": url, "status": pages[url], "location": "https://www.example.com/a" if pages[url] == 301 else None}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake_fetch)
        out = await gsc_sitemap_diff(SITE, "https://www.example.com/sitemap.xml",
                                     urls=["https://www.example.com/a", "https://www.example.com/live-only"])
        assert out["ok"] is True and out["inconclusive"] is False
        assert out["sitemap_url_count"] == 3  # image:loc ignored, index cycle handled
        assert out["missing_from_sitemap"] == ["https://www.example.com/live-only"]
        assert [r["url"] for r in out["sitemap_urls_not_200"]["redirect"]] == ["https://www.example.com/moved"]
        assert [r["url"] for r in out["sitemap_urls_not_200"]["not_found"]] == ["https://www.example.com/dead"]

    async def test_failed_control_marks_inconclusive(self, monkeypatch):
        async def refuse(*a, **k):
            return {"url": "x", "status": 403, "location": None, "body": b"<urlset/>"}
        monkeypatch.setattr(gsc_server, "_safe_fetch", refuse)
        out = await gsc_sitemap_diff(SITE, "https://www.example.com/sitemap.xml", urls=[])
        assert out["inconclusive"] is True and out["meta"]["warnings"]


# ---------------------------------------------------------------------------
# UI export import (acceptance test 8's mechanism)
# ---------------------------------------------------------------------------


def _compare_zip(path):
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("Pages.csv", "Top pages,8/1/26 - 8/31/26 Clicks,9/1/26 - 9/28/26 Clicks,"
                                 "8/1/26 - 8/31/26 Impressions,9/1/26 - 9/28/26 Impressions\n"
                                 "https://example.com/a,10,5,\"1,200\",300\n"
                                 "https://example.com/b,3,4,100,400\n")
        zf.writestr("Queries.csv", "Top queries,Clicks,Impressions,CTR,Position\nkw,5,100,5%,3.2\n")
        zf.writestr("Filters.csv", "Filter,Value\nSearch type,Web\nDate,Compare: Aug 2026 vs Sep 2026\n")


class TestUiExport:
    async def test_load_and_query_compare_export(self, tmp_path):
        path = tmp_path / "ai-features.zip"
        _compare_zip(path)
        loaded = await gsc_load_ui_export(str(path), report_kind="generative_ai_features")
        assert loaded["ok"] is True
        tables = {t["name"]: t for t in loaded["tables"]}
        assert {"pages", "queries", "filters", "meta"} <= set(tables)
        assert {"impr_aug", "impr_sep", "clicks_a", "clicks_b"} <= set(tables["pages"]["columns"])
        assert loaded["export_meta"]["search_type"] == "Web"
        out = await gsc_query_ui_export(
            loaded["session_id"],
            "select top_pages as page, impr_aug, impr_sep from pages order by (impr_aug - impr_sep) desc limit 20",
        )
        assert out["ok"] is True
        assert out["rows"][0] == {"page": "https://example.com/a", "impr_aug": 1200, "impr_sep": 300}
        q = await gsc_query_ui_export(loaded["session_id"], "select ctr, position from queries")
        assert q["rows"][0] == {"ctr": 0.05, "position": 3.2}
        w = await gsc_query_ui_export(
            loaded["session_id"],
            "select top_pages, rank() over (order by impr_sep desc) r, lag(impr_sep) over (order by top_pages) l from pages",
        )
        assert w["ok"] is True and len(w["rows"]) == 2
        m = await gsc_query_ui_export(loaded["session_id"], "select value from meta where key='date_range'")
        assert m["rows"][0]["value"].startswith("Compare")

    @pytest.mark.parametrize("sql", [
        "attach database '/tmp/x.db' as x",
        "pragma table_info(pages)",
        "insert into pages (top_pages) values ('x')",
        "select load_extension('x')",
        "create table t (a)",
        "select 1; select 2",
    ])
    async def test_non_read_sql_refused(self, tmp_path, sql):
        path = tmp_path / "e.zip"
        _compare_zip(path)
        sid = (await gsc_load_ui_export(str(path)))["session_id"]
        out = await gsc_query_ui_export(sid, sql)
        assert out["ok"] is False

    async def test_temp_store_is_memory(self, tmp_path):
        path = tmp_path / "e.zip"
        _compare_zip(path)
        sid = (await gsc_load_ui_export(str(path)))["session_id"]
        conn = gsc_server._ui_sessions[sid]["conn"]
        conn.set_authorizer(None)
        assert conn.execute("PRAGMA temp_store").fetchone()[0] == 2  # MEMORY
        conn.set_authorizer(gsc_server._ui_authorizer)

    async def test_limits(self, tmp_path, monkeypatch):
        big = tmp_path / "big.csv"
        big.write_text("a\n" + "\n".join(str(i) for i in range(6000)))
        out = await gsc_load_ui_export(str(big))
        assert out["ok"] is False and "5000" in out["error"]
        monkeypatch.setattr(gsc_server, "_UI_MAX_ZIP_ENTRIES", 1)
        z = tmp_path / "many.zip"
        _compare_zip(z)
        out = await gsc_load_ui_export(str(z))
        assert out["ok"] is False and "entries" in out["error"]
        assert (await gsc_load_ui_export("relative.csv"))["ok"] is False

    async def test_xlsx(self, tmp_path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Queries"
        ws.append(["Top queries", "Clicks", "Impressions"])
        ws.append(["kw", 3, 90])
        path = tmp_path / "export.xlsx"
        wb.save(path)
        loaded = await gsc_load_ui_export(str(path))
        out = await gsc_query_ui_export(loaded["session_id"], "select top_queries, impressions from queries")
        assert out["rows"] == [{"top_queries": "kw", "impressions": 90}]

    def test_sqlite_version_guard(self, monkeypatch):
        monkeypatch.setattr(sqlite3, "sqlite_version_info", (3, 24, 0))
        with pytest.raises(gsc_server._UiExportError):
            gsc_server._check_sqlite_version()


# ---------------------------------------------------------------------------
# Review fixes (open-gem pass 1 on Release B)
# ---------------------------------------------------------------------------


class TestReviewFixesB:
    async def test_fetch_is_pinned_to_validated_address(self, monkeypatch):
        monkeypatch.setattr(gsc_server.socket, "getaddrinfo",
                            lambda *a, **k: [(None, None, None, None, ("93.184.216.34", 443))])
        seen = {}

        class _Resp:
            status_code = 200
            headers = {}
            extensions = {}

        class _Ctx:
            async def __aenter__(self):
                return _Resp()

            async def __aexit__(self, *a):
                return False

        class _Client:
            def stream(self, method, url, headers=None, extensions=None):
                seen.update(url=url, host=headers["Host"], sni=extensions["sni_hostname"])
                return _Ctx()
        out = await gsc_server._safe_fetch(_Client(), "https://www.example.com/a?b=1", SITE)
        assert out["status"] == 200
        assert seen == {"url": "https://93.184.216.34/a?b=1", "host": "www.example.com", "sni": "www.example.com"}

    def test_sitemap_parsing_namespace_free_and_gzip_bomb(self):
        import gzip
        tag, locs = gsc_server._parse_sitemap_body(b"<urlset><url><loc>https://x.test/a</loc></url></urlset>")
        assert (tag, locs) == ("urlset", ["https://x.test/a"])
        bomb = gzip.compress(b"<urlset>" + b" " * (gsc_server._FETCH_MAX_BYTES + 10) + b"</urlset>")
        with pytest.raises(gsc_server._FetchRefused):
            gsc_server._parse_sitemap_body(bomb)
        with pytest.raises(ValueError):
            gsc_server._parse_sitemap_body(b"<html><body/></html>")

    async def test_incomplete_sitemap_only_possibly_missing(self, monkeypatch):
        index = b"""<sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
          <sitemap><loc>https://www.example.com/broken.xml</loc></sitemap></sitemapindex>"""

        async def fake_fetch(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": index}
            if url.endswith("broken.xml"):
                return {"url": url, "status": 500, "location": None, "body": b""}
            return {"url": url, "status": 200, "location": None}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake_fetch)
        out = await gsc_sitemap_diff(SITE, "https://www.example.com/sitemap.xml", urls=["https://www.example.com/a"])
        assert out["sitemap_complete"] is False
        assert out["missing_from_sitemap"] == [] and out["possibly_missing_from_sitemap"] == ["https://www.example.com/a"]

    def test_spike_uses_median_of_nonzero_days(self):
        rows = [{"query": "sparse", "clicks": 1, "impressions": 70}, {"query": "burst", "clicks": 0, "impressions": 600}]
        flags = gsc_server._machine_query_flags(rows, 10_000, {"sparse": [60, 55, 0, 0, 0], "burst": [5, 4, 500, 6]})
        assert "sparse" not in flags
        assert flags["burst"] == ["spike"]

    async def test_empty_terms_disable_config(self, monkeypatch):
        _write_config({SITE: {"brand_terms": ["example"]}})
        _patch(monkeypatch, _service(lambda b: {"rows": [_row("example", 1, 10)]} if b["dimensions"] else {"rows": [_row([], 1, 10)]}))
        out = await gsc_page_query_profile(SITE, "https://example.com/p", brand_terms=[])
        assert out["summary"]["brand_share"] is None

    async def test_ui_dimension_values_stay_text(self, tmp_path):
        p = tmp_path / "q.csv"
        p.write_text("Top queries,Clicks\n404,3\n2026 budget,1\n")
        sid = (await gsc_load_ui_export(str(p)))["session_id"]
        out = await gsc_query_ui_export(sid, "select top_queries, typeof(top_queries) t, clicks from q order by clicks desc")
        assert out["rows"][0] == {"top_queries": "404", "t": "text", "clicks": 3}

    async def test_ui_result_byte_budget(self, tmp_path, monkeypatch):
        p = tmp_path / "q.csv"
        p.write_text("Top queries,Clicks\n" + "\n".join(f"q{i},{i}" for i in range(100)))
        sid = (await gsc_load_ui_export(str(p)))["session_id"]
        monkeypatch.setattr(gsc_server, "_UI_MAX_RESULT_BYTES", 5000)
        out = await gsc_query_ui_export(sid, "select top_queries, substr(printf('%.*c', 2000, 'x'), 1, 2000) pad from q", limit=100)
        assert out["truncated"] is True and len(out["rows"]) < 100

    async def test_movers_net_from_full_totals(self, monkeypatch):
        calls = {"n": 0}

        def respond(body):
            calls["n"] += 1
            if calls["n"] == 1:
                return {"rows": [_row("/big", 100, 1000), _row("/tiny", 1, 5)]}
            return {"rows": [_row("/big", 50, 900), _row("/tiny", 11, 6)]}
        _patch(monkeypatch, _service(respond))
        out = await gsc_movers(SITE, days=7, min_impressions=100)
        assert out["net_change"]["clicks"] == -40  # includes /tiny although it is filtered out
        assert out["losers"][0]["share_of_clicks_change"] == pytest.approx(-50 / -40)


class TestWorklistReviewFixes:
    def _setup(self, monkeypatch, by_url, control_ok=True, live_status=200):
        TestRecrawlWorklist()._setup(monkeypatch, by_url)

        async def fake_control(client, site_url, sample):
            return {"url": "https://example.com/", "ok": control_ok, "status": 200 if control_ok else 403}

        async def fake_statuses(client, urls, site_url, **kw):
            return {u: {"url": u, "status": live_status} for u in urls}
        monkeypatch.setattr(gsc_server, "_control_ok", fake_control)
        monkeypatch.setattr(gsc_server, "_fetch_statuses", fake_statuses)

    async def _complete(self, items, **kw):
        for _ in range(200):
            out = await gsc_recrawl_worklist(SITE, items, **kw)
            if out["status"] == "complete":
                return out
            await asyncio.sleep(0.01)
        raise AssertionError("worklist never completed")

    async def test_redirect_with_failed_control_is_unresolved(self, monkeypatch):
        by_url = {"https://example.com/r": _inspection("NEUTRAL", "Page with redirect", "2026-09-20T00:00:00Z")}
        self._setup(monkeypatch, by_url, control_ok=False)
        out = await self._complete([{"url": "https://example.com/r", "changed_on": "2026-09-14"}])
        assert out["counts"]["unresolved"] == 1 and "control request failed" in out["unresolved"][0]["error"]["error"]

    async def test_confirmed_redirect_needs_nothing(self, monkeypatch):
        by_url = {"https://example.com/r": _inspection("NEUTRAL", "Page with redirect", "2026-09-20T00:00:00Z")}
        self._setup(monkeypatch, by_url, live_status=301)
        out = await self._complete([{"url": "https://example.com/r", "changed_on": "2026-09-14"}])
        assert out["counts"]["no_action_needed"] == 1 and "redirect is real" in out["no_action_needed"][0]["note"]

    async def test_retry_reinspects_only_failed_urls(self, monkeypatch):
        by_url = {"https://example.com/a": _inspection("PASS", "ok", "2026-09-20T00:00:00Z"),
                  "https://example.com/b": _inspection("PASS", "ok", "2026-09-20T00:00:00Z")}
        self._setup(monkeypatch, by_url)
        orig = gsc_server._inspect_one
        fail = {"b": True}

        async def flaky(service, site, url, **kw):
            if url.endswith("/b") and fail["b"]:
                raise RuntimeError("transient")
            return await orig(service, site, url, **kw)
        monkeypatch.setattr(gsc_server, "_inspect_one", flaky)
        items = [{"url": u, "changed_on": "2026-09-14"} for u in by_url]
        out = await self._complete(items)
        assert [u["url"] for u in out["unresolved"]] == ["https://example.com/b"]
        out = await self._complete(items)  # without retry: still reported, no new job
        assert out["counts"]["unresolved"] == 1
        used = gsc_server._inspection_quota(SITE)["day_used"]
        fail["b"] = False
        out = await self._complete(items, retry=True)
        assert out["counts"]["unresolved"] == 0
        assert gsc_server._inspection_quota(SITE)["day_used"] == used + 1  # only /b re-inspected


class TestReviewFixesB2:
    async def test_overlapping_worklists_share_active_jobs(self, monkeypatch):
        import threading
        gate = threading.Event()
        by_url = {f"https://example.com/{c}": _inspection("PASS", "ok", "2026-09-20T00:00:00Z") for c in "abc"}
        TestRecrawlWorklist()._setup(monkeypatch, by_url)
        orig = gsc_server._inspect_one

        async def slow(service, site, url, **kw):
            await asyncio.to_thread(gate.wait, 5)
            return await orig(service, site, url, **kw)
        monkeypatch.setattr(gsc_server, "_inspect_one", slow)
        first = await gsc_recrawl_worklist(SITE, [{"url": u, "changed_on": "2026-09-14"} for u in list(by_url)[:2]])
        second = await gsc_recrawl_worklist(SITE, [{"url": u, "changed_on": "2026-09-14"} for u in list(by_url)[1:]])
        assert first["status"] == second["status"] == "pending"
        # /b is covered by the first job; only /c starts a second one.
        assert len(gsc_server._inspect_jobs) == 2
        assert sorted(len(j["urls"]) for j in gsc_server._inspect_jobs.values()) == [1, 2]
        gate.set()

    def test_stale_job_result_not_reused(self):
        now = gsc_server.time.time()
        gsc_server._inspect_jobs["j"] = {
            "job_id": "j", "site_url": SITE, "urls": ["https://example.com/a"], "state": "done",
            "created_at": now, "results": {"https://example.com/a": {"url": "x"}},
            "result_ts": {"https://example.com/a": now - 7 * 3600}, "errors": {},
        }
        fresh, active, error = gsc_server._job_view(SITE, "https://example.com/a")
        assert fresh is None and active is None

    async def test_csv_row_cap_is_exact(self, tmp_path):
        ok = tmp_path / "ok.csv"
        ok.write_text("a\n" + "\n".join("x" for _ in range(5000)))
        assert (await gsc_load_ui_export(str(ok)))["ok"] is True
        bad = tmp_path / "bad.csv"
        bad.write_text("a\n" + "\n".join("x" for _ in range(5001)))
        assert (await gsc_load_ui_export(str(bad)))["ok"] is False

    async def test_control_rejects_external_redirect(self, monkeypatch):
        async def fake(client, url, site_url, **kw):
            return {"url": url, "status": 302, "location": "https://evil.test/"}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_server._control_ok(None, SITE, "https://www.example.com/x")
        assert out["ok"] is False
        async def fake_in(client, url, site_url, **kw):
            return {"url": url, "status": 301, "location": "/home"}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake_in)
        assert (await gsc_server._control_ok(None, SITE, "https://www.example.com/x"))["ok"] is True

    async def test_cannibalisation_save_to_file(self, monkeypatch, tmp_path):
        rows = []
        for i in range(30):
            rows += [_row([f"q{i}", "https://example.com/x"], 1, 50), _row([f"q{i}", "https://example.com/y"], 1, 50)]
        _patch(monkeypatch, _service(lambda b: {"rows": rows}))
        path = tmp_path / "c.json"
        out = await gsc_cannibalisation(SITE, limit=5, save_to_file=str(path))
        assert len(out["queries"]) == 20 and out["meta"]["saved_counts"]["queries"] == 30
        assert len(json.loads(path.read_text())["queries"]) == 30


class TestUiExportRealLayout:
    """Mirrors the real Generative AI features export (2026-09-28): day-first
    ranges in the headers, impressions only, no Queries/Dates tabs, Filters
    recording only the first range."""

    async def test_day_first_compare_xlsx(self, tmp_path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pages"
        ws.append(["Top pages", "01/08/2026 - 27/08/2026 Impressions", "01/09/2026 - 27/09/2026 Impressions"])
        ws.append(["https://example.com/a", 5999.0, 1239.0])
        ws.append(["https://example.com/b", 10.0, 20.0])
        f = wb.create_sheet("Filters")
        f.append(["Filter", "Value"])
        f.append(["Date", "1 Aug 2026-27 Aug 2026"])
        path = tmp_path / "ai.xlsx"
        wb.save(path)
        loaded = await gsc_load_ui_export(str(path))
        periods = loaded["tables"][0]["compare_periods"]
        assert [(p["start"], p["end"], p["month"]) for p in periods] == [
            ("2026-08-01", "2026-08-27", "aug"), ("2026-09-01", "2026-09-27", "sep")]
        assert loaded["export_meta"]["period_b"] == "2026-09-01..2026-09-27"
        out = await gsc_query_ui_export(loaded["session_id"],
                                        "select top_pages, impr_aug, impr_sep from pages order by impr_aug - impr_sep desc")
        assert out["rows"][0] == {"top_pages": "https://example.com/a", "impr_aug": 5999, "impr_sep": 1239}

    def test_ambiguous_dates_get_no_month_alias(self):
        names, aliases, periods = gsc_server._ui_columns(
            ["Top pages", "01/02/2026 - 07/02/2026 Clicks", "01/03/2026 - 07/03/2026 Clicks"])
        assert [p["month"] for p in periods] == [None, None]
        assert {a for a, _ in aliases} == {"clicks_a", "clicks_b"}


class TestFetchRateLimits:
    async def test_429_retried_then_reported_unchecked(self, monkeypatch):
        calls = {}

        async def fake(client, url, site_url, **kw):
            calls[url] = calls.get(url, 0) + 1
            if url.endswith("/flaky") and calls[url] < 3:
                return {"url": url, "status": 429, "location": None, "retry_after": "1"}
            if url.endswith("/always"):
                return {"url": url, "status": 429, "location": None, "retry_after": None}
            return {"url": url, "status": 200, "location": None}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_server._fetch_statuses(None, ["https://example.com/flaky", "https://example.com/always"], SITE)
        assert out["https://example.com/flaky"]["status"] == 200 and calls["https://example.com/flaky"] == 3
        assert out["https://example.com/always"]["rate_limited"] is True
        assert calls["https://example.com/always"] == gsc_server._FETCH_MAX_ATTEMPTS

    async def test_sitemap_diff_separates_rate_limited(self, monkeypatch):
        body = b"""<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
          <url><loc>https://www.example.com/ok</loc></url><url><loc>https://www.example.com/busy</loc></url></urlset>"""

        async def fake(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": body}
            if url.endswith("/busy"):
                return {"url": url, "status": 429, "location": None}
            return {"url": url, "status": 200, "location": None}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_sitemap_diff(SITE, "https://www.example.com/sitemap.xml", urls=[])
        assert out["sitemap_urls_unchecked_rate_limited"] == ["https://www.example.com/busy"]
        assert out["sitemap_urls_not_200"]["other"] == []
        assert out["checked"]["sitemap_urls"] == 1 and any("rate-limited" in w for w in out["meta"]["warnings"])


class TestBotChallenge:
    def _client(self, status, body, seen):
        class _Resp:
            status_code = status
            headers = {}
            extensions = {}

            async def aiter_bytes(self):
                yield body

        class _Ctx:
            async def __aenter__(self):
                return _Resp()

            async def __aexit__(self, *a):
                return False

        class _Client:
            def stream(self, method, url, headers=None, extensions=None):
                seen.append(url)
                return _Ctx()
        return _Client()

    async def test_challenge_page_flagged_and_not_retried(self, monkeypatch):
        monkeypatch.setattr(gsc_server.socket, "getaddrinfo", lambda *a, **k: [
            (None, None, None, None, ("2606:4700::1", 443, 0, 0)), (None, None, None, None, ("93.184.216.34", 443))])
        seen = []
        client = self._client(429, b"<html><title>Checking browser</title></html>", seen)
        r = await gsc_server._fetch_with_backoff(client, "https://www.example.com/p", SITE)
        assert r["bot_challenge"] is True and r["rate_limited"] is True
        assert len(seen) == 1  # no pointless retries against a challenge
        assert seen[0].startswith("https://[2606:4700::1]/")  # OS order kept: IPv6 first

    async def test_plain_429_still_retried(self, monkeypatch):
        monkeypatch.setattr(gsc_server.socket, "getaddrinfo",
                            lambda *a, **k: [(None, None, None, None, ("93.184.216.34", 443))])
        seen = []
        client = self._client(429, b"slow down", seen)
        r = await gsc_server._fetch_with_backoff(client, "https://www.example.com/p", SITE)
        assert r.get("bot_challenge") is None and r["rate_limited"] is True
        assert len(seen) == gsc_server._FETCH_MAX_ATTEMPTS


# ---------------------------------------------------------------------------
# v1.6.0: coverage on every sampled / capped result; redirect classification
# ---------------------------------------------------------------------------


class TestCoverage:
    def test_coverage_helper(self):
        c = gsc_server._coverage(40, 591, "sitemap URLs")
        assert c["partial"] is True and c["summary"] == "checked 40 of 591 sitemap URLs — PARTIAL"
        assert gsc_server._coverage(5, 5, "rows")["partial"] is False
        assert gsc_server._coverage(5, None, "rows")["partial"] is True

    async def test_query_coverage_partial_when_truncated(self, monkeypatch):
        _patch(monkeypatch, _service(lambda b: {"rows": [_row(f"q{i}", 1, 1) for i in range(b["rowLimit"])]}))
        out = await gsc_server.gsc_query(SITE, "2026-09-01", "2026-09-25", row_limit=10, include_totals=False)
        assert out["meta"]["coverage"]["partial"] is True and out["meta"]["coverage"]["of"] is None
        md = await gsc_server.gsc_query(SITE, "2026-09-01", "2026-09-25", row_limit=10, include_totals=False,
                                        response_format="markdown")
        assert "Coverage: checked 10 of an unknown total" in md

    async def test_query_coverage_complete(self, monkeypatch):
        _patch(monkeypatch, _service(lambda b: {"rows": [_row("q", 1, 1)]}))
        out = await gsc_server.gsc_query(SITE, "2026-09-01", "2026-09-25", include_totals=False)
        assert out["meta"]["coverage"] == {"checked": 1, "of": 1, "unit": "rows", "partial": False,
                                           "summary": "checked 1 of 1 rows"}

    async def test_cannibalisation_shown_vs_total(self, monkeypatch):
        rows = []
        for i in range(5):
            rows += [_row([f"q{i}", "https://example.com/x"], 1, 50), _row([f"q{i}", "https://example.com/y"], 1, 50)]
        _patch(monkeypatch, _service(lambda b: {"rows": rows}))
        out = await gsc_cannibalisation(SITE, limit=2)
        shown = out["meta"]["coverage"]["queries_shown"]
        assert (shown["checked"], shown["of"], shown["partial"]) == (2, 5, True)


class TestSitemapRedirectClassification:
    async def test_equivalent_hub_chain_and_cms_state(self, monkeypatch):
        W = "https://www.example.com"
        body = ("<urlset xmlns=\"http://www.sitemaps.org/schemas/sitemap/0.9\">" + "".join(
            f"<url><loc>{W}/{p}</loc></url>" for p in ("old-post", "old-news", "hop1", "loop", "fine", "unchecked")
        ) + "</urlset>").encode()
        pages = {
            f"{W}/": (200, None), f"{W}/fine": (200, None),
            f"{W}/old-post": (301, "/new-post"), f"{W}/new-post": (200, None),
            f"{W}/old-news": (301, f"{W}/blog/"), f"{W}/blog/": (200, None),
            f"{W}/hop1": (301, "/hop2"), f"{W}/hop2": (302, "/final"), f"{W}/final": (200, None),
            f"{W}/loop": (301, "/loop2"), f"{W}/loop2": (301, "/loop"),
            f"{W}/unchecked": (200, None),
        }

        async def fake(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": body}
            status, loc = pages[url]
            return {"url": url, "status": status, "location": loc}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_sitemap_diff(SITE, f"{W}/sitemap.xml", urls=[], check_limit=5,
                                     cms_state={f"{W}/old-post": "PUBLISHED", f"{W}/old-news": "archived"})
        by = {r["url"].rsplit("/", 1)[-1]: r for r in out["sitemap_urls_not_200"]["redirect"]}
        assert by["old-post"]["classification"] == "equivalent" and by["old-post"]["final_url"] == f"{W}/new-post"
        assert by["old-news"]["classification"] == "hub"
        assert by["hop1"]["classification"] == "chain" and by["hop1"]["hops"] == 2
        assert by["loop"]["classification"] == "chain" and by["loop"]["note"] == "redirect loop"
        assert by["old-post"]["published_but_redirected"] is True and by["old-news"]["published_but_redirected"] is False
        assert out["redirect_summary"] == {"equivalent": 1, "hub": 1, "chain": 2, "unverified": 0, "published_but_redirected": 1}
        # 5 of 6 sitemap URLs checked: partial, and said so up front.
        assert out["partial"] is True
        assert out["coverage"]["sitemap_urls_live_checked"]["summary"] == "checked 5 of 6 sitemap URLs live-checked — PARTIAL"
        assert any("continue with check_offset=5" in w for w in out["meta"]["warnings"])

    async def test_hub_paths_from_site_config(self, monkeypatch):
        W = "https://www.example.com"
        _write_config({SITE: {"hub_paths": ["/news"], "hub_regex": "^/tag/"}})
        body = f"<urlset xmlns=\"http://www.sitemaps.org/schemas/sitemap/0.9\"><url><loc>{W}/a</loc></url><url><loc>{W}/b</loc></url></urlset>".encode()
        pages = {f"{W}/": (200, None), f"{W}/a": (301, "/news/"), f"{W}/news/": (200, None),
                 f"{W}/b": (301, "/tag/x"), f"{W}/tag/x": (200, None)}

        async def fake(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": body}
            status, loc = pages[url]
            return {"url": url, "status": status, "location": loc}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_sitemap_diff(SITE, f"{W}/sitemap.xml", urls=[])
        assert {r["classification"] for r in out["sitemap_urls_not_200"]["redirect"]} == {"hub"}
        assert out["partial"] is False


class TestReviewFixesV16:
    async def test_unverified_destination_and_refused_not_counted(self, monkeypatch):
        W = "https://www.example.com"
        body = (f"<urlset xmlns=\"http://www.sitemaps.org/schemas/sitemap/0.9\"><url><loc>{W}/a</loc></url>"
                f"<url><loc>{W}/b</loc></url></urlset>").encode()

        async def fake(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": body}
            if url == f"{W}/":
                return {"url": url, "status": 200, "location": None}
            if url == f"{W}/a":
                return {"url": url, "status": 301, "location": "/busy"}
            if url == f"{W}/busy":
                return {"url": url, "status": 429, "location": None}
            raise gsc_server._FetchRefused("outside")
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_sitemap_diff(SITE, f"{W}/sitemap.xml", urls=[])
        red = out["sitemap_urls_not_200"]["redirect"][0]
        assert red["classification"] == "unverified"
        cov = out["coverage"]
        assert cov["sitemap_urls_live_checked"]["checked"] == 1  # /b was refused, not checked
        assert cov["redirects_classified"]["partial"] is True and out["partial"] is True

    def test_hub_paths_extend_defaults(self):
        hubs = {"/", "/blog", "/events", "/about", "/news"}
        info = {"hops": 1, "final_url": "https://x.test/blog/", "target_status": 200}
        assert gsc_server._classify_redirect(info, hubs, None) == "hub"

    async def test_configured_hubs_keep_defaults(self, monkeypatch):
        W = "https://www.example.com"
        _write_config({SITE: {"hub_paths": ["/news"]}})
        body = f"<urlset xmlns=\"http://www.sitemaps.org/schemas/sitemap/0.9\"><url><loc>{W}/a</loc></url></urlset>".encode()
        pages = {f"{W}/": (200, None), f"{W}/a": (301, "/blog"), f"{W}/blog": (200, None)}

        async def fake(client, url, site_url, *, extra_hosts=(), want_body=False, max_bytes=0):
            if url.endswith("sitemap.xml"):
                return {"url": url, "status": 200, "location": None, "body": body}
            status, loc = pages[url]
            return {"url": url, "status": status, "location": loc}
        monkeypatch.setattr(gsc_server, "_safe_fetch", fake)
        out = await gsc_sitemap_diff(SITE, f"{W}/sitemap.xml", urls=[])
        assert out["sitemap_urls_not_200"]["redirect"][0]["classification"] == "hub"
        assert "/news" in out["meta"]["hub_paths"] and "/blog" in out["meta"]["hub_paths"]
